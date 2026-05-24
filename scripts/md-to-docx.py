#!/usr/bin/env python3
"""Convert docs/THESIS_CH4.md to Word (.docx) with embedded screenshots."""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
except ImportError:
    print("Installing python-docx...")
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-q"])
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
MD_PATH = ROOT / "docs" / "THESIS_CH4.md"
OUT_PATH = ROOT / "docs" / "THESIS_CH4.docx"
OUT_PATH_ALT = ROOT / "docs" / "THESIS_CH4_updated.docx"

# Thai thesis fonts (Windows usually has TH Sarabun New or Leelawadee UI)
BODY_FONT = "TH Sarabun New"
FALLBACK_FONT = "Leelawadee UI"
CODE_FONT = "Consolas"


def set_run_font(run, name: str, size_pt: float, bold: bool = False, color: RGBColor | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color


def add_formatted_paragraph(doc: Document, text: str, style: str = "Normal", size: float = 16):
    """Parse inline **bold** and `code` in a paragraph."""
    p = doc.add_paragraph(style=style)
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, CODE_FONT, size - 1)
            run.font.color.rgb = RGBColor(0x80, 0x00, 0x00)
        else:
            run = p.add_run(part)
            set_run_font(run, BODY_FONT, size)


def add_code_block(doc: Document, lines: list[str]):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    text = "\n".join(lines)
    run = p.add_run(text)
    set_run_font(run, CODE_FONT, 11)
    run.font.color.rgb = RGBColor(0x1E, 0x1E, 0x1E)
    # light gray shading via paragraph
    from docx.oxml import OxmlElement

    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F5F5F5")
    p._p.get_or_add_pPr().append(shd)


def parse_table_row(line: str) -> list[str]:
    line = line.strip().strip("|")
    return [c.strip() for c in line.split("|")]


def is_table_separator(line: str) -> bool:
    return bool(re.match(r"^\|?[\s\-:|]+\|?$", line.strip()))


def add_table(doc: Document, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            if j >= cols:
                break
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(cell_text.replace("`", ""))
            set_run_font(run, BODY_FONT, 14, bold=(i == 0))


def add_image(doc: Document, img_path: Path, caption: str = ""):
    if not img_path.exists():
        add_formatted_paragraph(doc, f"[ไม่พบรูป: {img_path.name}]", size=14)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    # fit to page width ~16cm
    run.add_picture(str(img_path), width=Cm(16))
    if caption:
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        set_run_font(r, BODY_FONT, 14, bold=True)
        r.italic = True


def convert(md_path: Path, out_path: Path):
    doc = Document()

    # Page margins (thesis style)
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    # Default style
    style = doc.styles["Normal"]
    style.font.name = BODY_FONT
    style.font.size = Pt(16)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    table_buffer: list[list[str]] = []
    in_code = False
    code_lines: list[str] = []
    code_lang = ""

    img_re = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")

    while i < len(lines):
        line = lines[i]

        # Code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lang = line.strip()[3:].strip()
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
                code_lines = []
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # Flush table
        if table_buffer and not line.strip().startswith("|"):
            add_table(doc, table_buffer)
            table_buffer = []

        # Table row
        if line.strip().startswith("|"):
            if is_table_separator(line):
                i += 1
                continue
            table_buffer.append(parse_table_row(line))
            i += 1
            continue

        # Horizontal rule
        if line.strip() in ("---", "***", "___"):
            doc.add_paragraph()
            i += 1
            continue

        # Image
        m = img_re.match(line.strip())
        if m:
            caption, rel = m.group(1), m.group(2)
            img_path = (md_path.parent / rel).resolve()
            add_image(doc, img_path, caption)
            i += 1
            continue

        # Headings
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=0)
            for run in p.runs:
                set_run_font(run, BODY_FONT, 22, bold=True)
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=1)
            for run in p.runs:
                set_run_font(run, BODY_FONT, 20, bold=True)
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=2)
            for run in p.runs:
                set_run_font(run, BODY_FONT, 18, bold=True)
            i += 1
            continue

        # Ordered list
        om = re.match(r"^(\d+)\.\s+(.*)$", line.strip())
        if om:
            p = doc.add_paragraph(style="List Number")
            add_formatted_paragraph_inline(p, om.group(2), size=16)
            i += 1
            continue

        # Bullet list
        if line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_paragraph_inline(p, line.strip()[2:], size=16)
            i += 1
            continue

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Italic footer
        if line.strip().startswith("*") and line.strip().endswith("*"):
            p = doc.add_paragraph()
            run = p.add_run(line.strip().strip("*"))
            set_run_font(run, BODY_FONT, 14)
            run.italic = True
            i += 1
            continue

        # Normal paragraph
        add_formatted_paragraph(doc, line.strip(), size=16)
        i += 1

    if table_buffer:
        add_table(doc, table_buffer)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        doc.save(str(out_path))
        print(f"Saved: {out_path}")
    except PermissionError:
        doc.save(str(OUT_PATH_ALT))
        print(f"Saved (alternate): {OUT_PATH_ALT}")
        print("Note: Close THESIS_CH4.docx if open in Word, then rename or re-run.")


def add_formatted_paragraph_inline(p, text: str, size: float = 16):
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, BODY_FONT, size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, CODE_FONT, size - 1)
        else:
            run = p.add_run(part)
            set_run_font(run, BODY_FONT, size)


if __name__ == "__main__":
    md = Path(sys.argv[1]) if len(sys.argv) > 1 else MD_PATH
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else OUT_PATH
    convert(md, out)
